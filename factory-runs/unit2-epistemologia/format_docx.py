from __future__ import annotations
import hashlib, json, pathlib, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if len(sys.argv) != 4:
    raise SystemExit("USAGE: format_docx.py INPUT.docx FORMATTING_EVIDENCE.json UNIACC_LOGO.png")
docx_path = pathlib.Path(sys.argv[1])
evidence_path = pathlib.Path(sys.argv[2])
logo_path = pathlib.Path(sys.argv[3])
if not docx_path.is_file() or not logo_path.is_file():
    raise SystemExit("INPUT_OR_LOGO_MISSING")

doc = Document(str(docx_path))
for section in doc.sections:
    section.top_margin = Inches(2.5 / 2.54)
    section.bottom_margin = Inches(2.5 / 2.54)
    section.left_margin = Inches(2.5 / 2.54)
    section.right_margin = Inches(2.5 / 2.54)
    section.header_distance = Inches(1.25 / 2.54)
    section.footer_distance = Inches(1.25 / 2.54)
    # APA 7 student-paper page number: right-aligned in the header on every page.
    header_paragraph = section.header.paragraphs[0]
    header_paragraph.clear()
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_paragraph.paragraph_format.space_after = Pt(0)
    page_run = header_paragraph.add_run()
    page_run.font.name = "Arial"
    page_run.font.size = Pt(12)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    field_instruction = OxmlElement("w:instrText")
    field_instruction.set(qn("xml:space"), "preserve")
    field_instruction.text = " PAGE "
    field_separator = OxmlElement("w:fldChar")
    field_separator.set(qn("w:fldCharType"), "separate")
    field_value = OxmlElement("w:t")
    field_value.text = "1"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    for node in (field_begin, field_instruction, field_separator, field_value, field_end):
        page_run._r.append(node)

for style_name in ("Normal", "Body Text"):
    if style_name in doc.styles:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
    if style_name in doc.styles:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(16 if style_name != "Subtitle" else 12)
        if style_name != "Subtitle":
            style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

phase = "COVER"
logo_inserted = False
page_headings = []
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt == "[[UNIACC_LOGO]]":
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.add_run().add_picture(str(logo_path), width=Inches(2.65))
        logo_inserted = True
        continue
    if txt.startswith("Pregunta 1.") or txt.startswith("Pregunta 2.") or txt.startswith("Pregunta 3.") or txt.startswith("Pregunta 4.") or txt == "Referencias":
        p.paragraph_format.page_break_before = True
        page_headings.append(txt)
        phase = "REFERENCES" if txt == "Referencias" else "BODY"
    if phase == "COVER":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        if txt.startswith("IDENTIFICACIÓN DE LOS ENFOQUES"):
            p.style = doc.styles["Title"]
        elif txt in ("PSICOLOGÍA", "EPISTEMOLOGÍA - 387 - 487971"):
            for run in p.runs:
                run.bold = True
    elif phase == "REFERENCES":
        if txt == "Referencias":
            # Do not assume the localized Heading 1 style exists in every Pandoc template.
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
    elif p.style.name.startswith("Heading"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        run.font.name = "Arial"
        if p.style.name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            run.font.size = Pt(16)
        else:
            run.font.size = Pt(12)

if not logo_inserted:
    raise SystemExit("UNIACC_LOGO_MARKER_NOT_FOUND")

doc.core_properties.title = "Identificación de los enfoques de la epistemología científica"
doc.core_properties.subject = "Psicología — Epistemología, Unidad 2"
doc.core_properties.author = "Diego Ignacio Norambuena Miranda"
doc.core_properties.keywords = "epistemología, positivismo lógico, Kuhn, Quine, Nersessian, APA 7"
doc.save(str(docx_path))
digest = hashlib.sha256(docx_path.read_bytes()).hexdigest()
evidence_path.parent.mkdir(parents=True, exist_ok=True)
evidence_path.write_text(json.dumps({
    "schema": "UNIT2_EPISTEMOLOGY_DOCX_FORMATTING_EVIDENCE/1.0",
    "status": "APPLIED",
    "artifact": docx_path.name,
    "sha256": digest,
    "font": "Arial",
    "body_size_pt": 12,
    "heading_size_pt": 16,
    "line_spacing": 1.5,
    "paragraph_alignment": "justified",
    "cover_alignment": "centered",
    "references_hanging_indent": True,
    "cover_logo": "UNIACC official logo retrieved from Wikimedia Commons",
    "page_numbers": "APA 7 PAGE field, right-aligned in header on each page",
    "logo_sha256": hashlib.sha256(logo_path.read_bytes()).hexdigest(),
    "page_break_before": page_headings,
    "margins_cm": {"top": 2.5, "bottom": 2.5, "left": 2.5, "right": 2.5},
    "author": "Diego Ignacio Norambuena Miranda",
    "g23": "NOT_EXECUTED",
    "g24": "NOT_EXECUTED"
}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status":"PASS","artifact":docx_path.name,"sha256":digest,"logo_inserted":logo_inserted,"page_breaks":len(page_headings)},ensure_ascii=False,sort_keys=True))
